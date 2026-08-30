import os
import io
import base64
import re
from bs4 import BeautifulSoup
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def hex_to_rgb(hex_str):
    hex_str = hex_str.lstrip('#')
    if len(hex_str) == 6:
        return RGBColor(int(hex_str[0:2], 16), int(hex_str[2:4], 16), int(hex_str[4:6], 16))
    return RGBColor(255, 255, 255)

def set_slide_background(slide, color):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_text_to_shape(tf, text, font_size, color, bold=False, align=PP_ALIGN.LEFT, space_after=6):
    p = tf.add_paragraph() if len(tf.paragraphs[0].text) > 0 else tf.paragraphs[0]
    p.alignment = align
    p.space_after = Pt(space_after)
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = "Calibri"
    return p

def parse_html_to_pptx(html_path, pptx_path):
    with open(html_path, 'r', encoding='utf-8') as f:
        soup = BeautifulSoup(f.read(), 'html.parser')

    prs = Presentation()
    # 16:9 Widescreen
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    bg_color = RGBColor(0, 11, 24)       # #000b18
    title_color = RGBColor(0, 210, 255)  # #00d2ff
    text_color = RGBColor(241, 245, 249) # #f1f5f9
    sub_color = RGBColor(203, 213, 225)  # #cbd5e1
    accent_color = RGBColor(0, 255, 136) # #00ff88

    sections = soup.find_all('section')
    print(f"Found {len(sections)} slides.")

    for idx, sec in enumerate(sections):
        slide = prs.slides.add_slide(blank_layout)
        set_slide_background(slide, bg_color)

        # Check for title
        title_tag = sec.find(['h1', 'h2'])
        title_text = title_tag.get_text().strip() if title_tag else ""
        
        if title_text:
            title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(1.0))
            tf = title_box.text_frame
            tf.word_wrap = True
            add_text_to_shape(tf, title_text, 32, title_color, bold=True, align=PP_ALIGN.CENTER if idx==0 else PP_ALIGN.LEFT)

        # Handle embedded images
        imgs = sec.find_all('img')
        img_offset_x = Inches(7.0)
        has_right_img = False
        
        for img in imgs:
            src = img.get('src', '')
            if src.startswith('data:image'):
                try:
                    b64_data = src.split(',')[1]
                    img_bytes = base64.b64decode(b64_data)
                    img_stream = io.BytesIO(img_bytes)
                    
                    if idx == 0:
                        # Center logo on title slide
                        slide.shapes.add_picture(img_stream, Inches(5.166), Inches(2.2), width=Inches(3.0))
                    else:
                        # Add image to right side
                        has_right_img = True
                        slide.shapes.add_picture(img_stream, Inches(6.8), Inches(1.8), width=Inches(5.8))
                except Exception as e:
                    print(f"Error decoding image on slide {idx+1}: {e}")

        # Find columns / content containers
        # Look for flex columns or cards
        cols = []
        for div in sec.find_all('div', recursive=True):
            style = div.get('style', '')
            if 'flex: 1' in style or 'flex:1' in style or 'col' in div.get('class', []):
                # Avoid nested duplicates
                if not any(div in c.descendants for c in cols):
                    cols.append(div)

        if not cols and idx > 0:
            # If no flex columns, use the main section body
            cols = [sec]

        # Determine widths and positions based on columns & images
        num_cols = len(cols) if len(cols) > 0 else 1
        max_width = Inches(5.6) if has_right_img else Inches(11.733)
        col_width = max_width / num_cols
        
        for c_idx, col in enumerate(cols):
            left_pos = Inches(0.8) + (c_idx * col_width)
            top_pos = Inches(1.6) if title_text else Inches(1.0)
            
            # Create a card background if it looks like a card
            style = col.get('style', '')
            if 'background' in style or 'border' in style or 'padding' in style:
                shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_pos + Inches(0.1), top_pos, col_width - Inches(0.2), Inches(5.2))
                shape.fill.solid()
                shape.fill.fore_color.rgb = RGBColor(10, 25, 47) # Dark blue card
                shape.line.color.rgb = RGBColor(30, 58, 95)
            
            tb = slide.shapes.add_textbox(left_pos + Inches(0.2), top_pos + Inches(0.2), col_width - Inches(0.4), Inches(4.8))
            tf = tb.text_frame
            tf.word_wrap = True

            # Process elements inside column
            elements = col.find_all(['h3', 'h4', 'p', 'li', 'div'], recursive=False)
            if not elements:
                elements = col.find_all(['h3', 'h4', 'p', 'li'])

            for el in elements:
                if el.name in ['h1', 'h2'] and el == title_tag:
                    continue
                
                text = el.get_text().strip()
                if not text:
                    continue

                if el.name == 'h3':
                    add_text_to_shape(tf, text, 22, title_color, bold=True, space_after=10)
                elif el.name == 'h4':
                    add_text_to_shape(tf, text, 18, accent_color, bold=True, space_after=8)
                elif el.name == 'li':
                    p = add_text_to_shape(tf, "• " + text, 15, text_color, space_after=6)
                elif el.name == 'p':
                    add_text_to_shape(tf, text, 15, sub_color, space_after=8)
                elif el.name == 'div':
                    # Could be an alert box or badge
                    add_text_to_shape(tf, text, 14, text_color, space_after=8)

        # Special handling for Slide 1 text below logo
        if idx == 0:
            tb = slide.shapes.add_textbox(Inches(1.5), Inches(5.4), Inches(10.333), Inches(1.5))
            tf = tb.text_frame
            tf.word_wrap = True
            add_text_to_shape(tf, "Un moteur NER (Named Entity Recognition) au service du médical", 20, text_color, bold=True, align=PP_ALIGN.CENTER)
            add_text_to_shape(tf, "L'ÉQUIPE JEDHA AIFS01\nPatrick Mouliom • Christopher Gilleron • Jérémie Becker • Arnaud Hoarau • Karim Atebata", 14, sub_color, align=PP_ALIGN.CENTER)

    prs.save(pptx_path)
    print(f"Successfully saved PowerPoint presentation to {pptx_path}")

if __name__ == '__main__':
    html_file = os.path.join(os.path.dirname(__file__), '..', 'docs', 'demoday_cliner.html')
    pptx_file = os.path.join(os.path.dirname(__file__), '..', 'docs', 'demoday_cliner.pptx')
    parse_html_to_pptx(html_file, pptx_file)
